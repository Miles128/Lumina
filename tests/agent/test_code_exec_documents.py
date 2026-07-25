"""Tests for code_exec and read_document tools."""

from __future__ import annotations

from pathlib import Path

from secretary.agent.agent_profile import AgentProfile, resolve_parent_tools
from secretary.agent.confirmation_policy import tool_requires_confirmation
from secretary.agent.tools.code_exec import CodeExecTool
from secretary.agent.tools.documents import ReadDocumentTool
from secretary.agent.tools.fs import FileReadTool, ListDirTool
from secretary.services.file_auth import FileAuthService


def test_code_exec_runs_snippet(tmp_path: Path) -> None:
    tool = CodeExecTool()
    out = tool.execute({"code": "print(2 + 2)"}, tmp_path)
    assert "4" in out
    assert tool.needs_confirmation is True


def test_code_exec_captures_stderr_and_exit(tmp_path: Path) -> None:
    tool = CodeExecTool()
    out = tool.execute({"code": "import sys\nprint('bad', file=sys.stderr)\nsys.exit(3)"}, tmp_path)
    assert "bad" in out
    assert "exit code: 3" in out


def test_code_exec_rejects_empty(tmp_path: Path) -> None:
    tool = CodeExecTool()
    assert tool.execute({"code": "  "}, tmp_path).startswith("Error:")


def test_code_exec_reads_workspace_absolute(tmp_path: Path) -> None:
    data = tmp_path / "numbers.txt"
    data.write_text("1\n2\n3\n", encoding="utf-8")
    tool = CodeExecTool()
    out = tool.execute(
        {
            "code": (
                f"print(sum(int(line) for line in open({str(data)!r})))\n"
            )
        },
        tmp_path,
    )
    assert "6" in out


def test_code_exec_reads_workspace_relative(tmp_path: Path) -> None:
    (tmp_path / "sample.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    tool = CodeExecTool()
    out = tool.execute(
        {"code": "print(open('sample.csv').read().strip())\n"},
        tmp_path,
    )
    assert "1,2" in out


def test_code_exec_blocks_write_to_workspace(tmp_path: Path) -> None:
    target = tmp_path / "owned.txt"
    target.write_text("keep", encoding="utf-8")
    tool = CodeExecTool()
    out = tool.execute(
        {"code": f"open({str(target)!r}, 'w').write('hacked')\n"},
        tmp_path,
    )
    assert target.read_text(encoding="utf-8") == "keep"
    assert "Error" in out or "PermissionError" in out or "sandbox" in out.lower()


def test_code_exec_allows_write_inside_sandbox(tmp_path: Path) -> None:
    tool = CodeExecTool()
    out = tool.execute(
        {
            "code": (
                "open('out.txt', 'w').write('ok')\n"
                "print(open('out.txt').read())\n"
            )
        },
        tmp_path,
    )
    assert "ok" in out
    assert "Error" not in out.split("[stderr]")[0]


def test_code_exec_blocks_open_outside_cwd(tmp_path: Path) -> None:
    secret = tmp_path / "secret.txt"
    secret.write_text("top-secret", encoding="utf-8")
    # Snippet runs with workspace=tmp_path so absolute under workspace is allowed.
    # Outside both sandbox and workspace must fail.
    outside = tmp_path.parent / "outside-secret.txt"
    outside.write_text("top-secret", encoding="utf-8")
    tool = CodeExecTool()
    out = tool.execute(
        {"code": f"print(open({str(outside)!r}).read())"},
        tmp_path,
    )
    assert "top-secret" not in out
    assert "Error" in out or "PermissionError" in out or "sandbox" in out.lower()


def test_code_exec_blocks_network_socket(tmp_path: Path) -> None:
    tool = CodeExecTool()
    out = tool.execute(
        {
            "code": (
                "import socket\n"
                "s = socket.socket()\n"
                "s.connect(('127.0.0.1', 1))\n"
            )
        },
        tmp_path,
    )
    assert "Error" in out or "PermissionError" in out or "sandbox" in out.lower()


def test_code_exec_session_grant_skips_confirmation(tmp_path: Path) -> None:
    tool = CodeExecTool()
    auth = FileAuthService(tmp_path / "file_auth.json")
    needs, kind = tool_requires_confirmation(
        tool, {"code": "print(1)"}, working_dir=tmp_path, file_auth=auth
    )
    assert needs is True
    assert kind == "action"
    auth.grant_session_code_exec()
    needs2, kind2 = tool_requires_confirmation(
        tool, {"code": "print(1)"}, working_dir=tmp_path, file_auth=auth
    )
    assert needs2 is False
    assert kind2 == ""


def test_ask_profile_includes_code_exec_plan_excludes(tmp_path: Path) -> None:
    tools = [ListDirTool(), FileReadTool(), CodeExecTool()]
    ask = {t.name for t in resolve_parent_tools(AgentProfile.ASK, tools, spawn_tool=None)}
    plan = {t.name for t in resolve_parent_tools(AgentProfile.PLAN, tools, spawn_tool=None)}
    assert "code_exec" in ask
    assert "code_exec" not in plan


def test_read_document_xlsx(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws.append(["name", "score"])
    ws.append(["Alice", 95])
    wb.save(path)

    tool = ReadDocumentTool()
    out = tool.execute({"path": str(path)}, tmp_path)
    assert "Alice" in out
    assert "95" in out
    assert "Data" in out


def test_read_document_docx(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "note.docx"
    doc = Document()
    doc.add_paragraph("Hello from Word")
    doc.save(path)

    tool = ReadDocumentTool()
    out = tool.execute({"path": str(path)}, tmp_path)
    assert "Hello from Word" in out


def test_read_document_pdf(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    path = tmp_path / "doc.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.write(path)

    tool = ReadDocumentTool()
    out = tool.execute({"path": str(path)}, tmp_path)
    assert "Error:" not in out
    assert "PDF" in out or "Page" in out


def test_read_document_unsupported(tmp_path: Path) -> None:
    path = tmp_path / "a.bin"
    path.write_bytes(b"\x00\x01")
    tool = ReadDocumentTool()
    out = tool.execute({"path": str(path)}, tmp_path)
    assert out.startswith("Error: unsupported")
